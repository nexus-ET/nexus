"""Quick assertions for ScanX subject+marks + transcript field parsing.

Run: python -m pytest backend/scripts/test_scanx_tamil_subjects.py -q
  or: python backend/scripts/test_scanx_tamil_subjects.py
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.services.scanx_academic_parse import (
    parse_stacked_marksheet_subjects,
    parse_subjects_from_text,
    resolve_subjects_from_text,
)
from app.services.scanx_categorize import categorize_extracted_text
from app.services.scanx_field_extract import extract_labeled_fields

# Real RapidOCR from scanx_documents.id=21 (nexus_edutrust_dev) — full footer.
REAL_TN_HSC_OCR = """
e
11180882
CERTIFICATE SL. NO. HSG
STATE BOARD OF SCHOOL EXAMINATIONS, TAMILNADU
DEPARTMENT OF GOVERNMENT EXAMINATIONS, CHENNAI - 600 006
HIGHER SECONDARY COURSE CERTIFICATE
GUITG6/GENERAL EDUCATION
ISSUED UNDER THE AUTHORITY OF THE GOVERNMENT OF TAMILNADU
SARAVANAVEL C
MAR 2016
and obtained the following marks :
SUBJECT
THORY
160
PRAC.
50
MARKS OBTAINED FOR 200
TAMIL
169
ONE SIX NINE
(P)
ENGLISH
137
ONE THREE SEVEN
(d)
PHYSICS
097
050
147
ONE FOUR SEVEN
(d)
CHEMISTRY
060
050
110
ONE ONE ZERO
(d)
BIOLOGY
056
050
106
ONE ZERO SIX
(d)
MATHEMATICS
155
ONE FIVE FIVE
(d)
GL LGLIT/TOTAL MARKS:
0824
ZERO EIGHT TWO FOUR
/DATE OF BIRTH
G/ ROLL NO.
.LD. & /MR CODE NO. DATE
06.06.1999
478100
G526345
17.05.2016
U/ PERMANENT REGISTER NO.
ug Q/MEDIUM OF INSTRUCTION
Lm_ e / GROUP CODE
1610468100
TAMIL
103
INFANT JESUS HR SEC SCHOOL S VINAYAGAPURAM TIRUCHIRAPPALLI
GU /NAME OF THE SCHOOL
MEMBER SECRETARY
STATE BOARD OF SCHOOL EXAMINATIONS (HR SEC), TAMLNADU
"""

SAMPLES = [
    ("Tamil 169", "space"),
    ("Tamil\t169", "tab"),
    ("Tamil | 169", "pipe"),
    ("Tamil: 169", "colon"),
    ("Tamil - 169", "hyphen"),
    ("Subject: Tamil\nMarks: 169", "labeled pair"),
    ("Subject : Tamil\nScore : 169", "score label"),
    ("Subject: Tamil\nMarks Obtained: 169", "marks obtained"),
    ("SUBJECT MARKS\nTAMIL 169\nENGLISH 145", "header block"),
    (REAL_TN_HSC_OCR, "real TN HSC stacked OCR"),
]


def _has_tamil_169(rows: list[dict]) -> bool:
    return any(
        (s.get("name") or "").strip().lower() == "tamil" and str(s.get("marks") or "") == "169"
        for s in rows
    )


def _field_map(text: str) -> dict[str, str]:
    return {f["label"]: f["value"] for f in extract_labeled_fields(text)}


def test_tamil_169_all_layouts() -> None:
    for text, label in SAMPLES:
        rows = parse_subjects_from_text(text)
        assert _has_tamil_169(rows), f"{label}: expected Tamil/169, got {rows}"


def test_stacked_parser_real_ocr() -> None:
    rows = parse_stacked_marksheet_subjects(REAL_TN_HSC_OCR)
    assert _has_tamil_169(rows), rows
    names = {(s.get("name") or "").strip().lower() for s in rows}
    assert "english" in names
    assert "physics" in names
    assert "mathematics" in names
    physics = next(s for s in rows if (s.get("name") or "").lower() == "physics")
    assert physics.get("marks") == "147"  # total, not theory 097


def test_resolve_merges_without_llm_when_stacked_ok() -> None:
    rows = resolve_subjects_from_text(REAL_TN_HSC_OCR, use_llm=False)
    assert _has_tamil_169(rows), rows
    assert len(rows) >= 5


def test_transcript_fields_dob_roll_group() -> None:
    got = _field_map(REAL_TN_HSC_OCR)
    assert got.get("Date of Birth") == "06.06.1999", got
    assert got.get("Roll Number") == "478100", got
    assert got.get("Group Code") == "103", got
    assert got.get("TMR Code No") == "G526345", got
    assert got.get("Issue Date") == "17.05.2016", got
    assert got.get("Permanent Register No") == "1610468100", got
    assert got.get("Medium of Instruction") == "TAMIL", got
    assert got.get("Candidate Name") == "SARAVANAVEL C", got
    assert "INFANT JESUS" in (got.get("School") or "")
    assert got.get("Total Marks") == "0824", got


def test_transcript_fields_after_normalize() -> None:
    """Persisted extracted_text is noise-cleaned; TMR code must survive."""
    from app.services.scanx_embeddings import normalize_extracted_display

    norm = normalize_extracted_display(REAL_TN_HSC_OCR)
    assert "G526345" in norm
    got = _field_map(norm)
    assert got.get("TMR Code No") == "G526345", got
    assert got.get("Total Marks") == "0824", got
    assert got.get("Issue Date") == "17.05.2016", got


def test_tmr_label_not_truncated_to_mr() -> None:
    text = REAL_TN_HSC_OCR.replace("/MR CODE NO. DATE", "/TMR CODE NO. DATE")
    got = _field_map(text)
    assert got.get("TMR Code No") == "G526345", got
    assert "MR Code No" not in got


def test_mathematics_survives_misplaced_total() -> None:
    """DOCX/OCR reading order may emit TOTAL MARKS before Mathematics."""
    text = """
SUBJECT
TAMIL
169
ONE SIX NINE
(P)
ENGLISH
137
ONE THREE SEVEN
(d)
GL LGLIT/TOTAL MARKS:
0824
ZERO EIGHT TWO FOUR
MATHEMATICS
155
ONE FIVE FIVE
(d)
/DATE OF BIRTH
"""
    rows = parse_stacked_marksheet_subjects(text)
    names = {(s.get("name") or "").strip().lower() for s in rows}
    assert "mathematics" in names, rows
    assert "tamil" in names
    math = next(s for s in rows if (s.get("name") or "").lower() == "mathematics")
    assert math.get("marks") == "155"

def test_categorizer_includes_tamil_subjects() -> None:
    payload = categorize_extracted_text(
        "Subject: Tamil\nMarks: 169\nEnglish 145",
        document_type_id="GRADE_SHEET",
        use_llm_fields=False,
    )
    assert payload is not None
    assert _has_tamil_169(list(payload.get("subjects") or []))


def test_categorizer_real_ocr() -> None:
    payload = categorize_extracted_text(
        REAL_TN_HSC_OCR,
        document_type_id="TR_TRANSCRIPT",
        use_llm_fields=False,
    )
    assert payload is not None
    assert _has_tamil_169(list(payload.get("subjects") or [])), payload.get("subjects")
    fields = {f["label"]: f["value"] for f in (payload.get("fields") or [])}
    assert fields.get("Date of Birth") == "06.06.1999", fields
    assert fields.get("Roll Number") == "478100", fields
    assert fields.get("Group Code") == "103", fields
    assert fields.get("TMR Code No") == "G526345", fields
    assert fields.get("Maximum Marks") == "200", fields
    assert fields.get("Total Marks") == "0824", fields

    # Heading banners must never appear as Score / Academic values.
    academic = next(c for c in payload["categories"] if c["id"] == "academic")
    values = [i["value"] for i in academic["items"]]
    labels = [i["label"] for i in academic["items"]]
    assert "MARKS OBTAINED FOR 200" not in values
    assert "GU /NAME OF THE SCHOOL" not in values
    assert "Score / grade" not in labels or all(
        i["label"] != "Score / grade" or "OBTAINED" not in i["value"]
        for i in academic["items"]
    )
    # No garbage subject crumbs.
    names = {(s.get("name") or "").lower() for s in (payload.get("subjects") or [])}
    assert "gu lg" not in names


def test_mr_code_split_fragments_merge() -> None:
    """OCR sometimes splits G526345 into G52 + 6345 — merge under TMR Code No."""
    text = REAL_TN_HSC_OCR.replace("G526345", "G52\n6345")
    got = _field_map(text)
    assert got.get("TMR Code No") == "G526345", got
    # Must not invent a heading from the fragment.
    assert "G52" not in got
    assert got.get("Issue Date") == "17.05.2016", got


def test_light_regex_clean_preserves_tmr() -> None:
    from app.services.ocr_cleaner import light_regex_clean

    assert "TMR" in light_regex_clean("TMR CODE NO.")
    assert light_regex_clean("###||| TAMIL 169") != ""


def test_gibberish_ocr_artifacts_dropped() -> None:
    from app.services.ocr_cleaner import is_ocr_gibberish_line, light_regex_clean
    from app.services.scanx_embeddings import clean_ocr_extracted_text, _ocr_line_is_noise

    for junk in ("Ciamr aGwrium", "Epulona Glowes", "ulona Gloweos"):
        assert is_ocr_gibberish_line(junk), junk
        assert _ocr_line_is_noise(junk), junk
        assert light_regex_clean(junk) == "", junk
    cleaned = clean_ocr_extracted_text(
        "TAMIL\n169\nCiamr aGwrium\nEpulona Glowes\nTMR CODE NO.\nG526345"
    )
    assert "Ciamr" not in cleaned
    assert "Epulona" not in cleaned
    assert "TAMIL" in cleaned
    assert "G526345" in cleaned


def test_theory_practical_total_fields() -> None:
    """Theory/Practical are per-subject columns — not Document field labels.

    Certificate maxima (THORY 160 / PRAC 50) fold into Maximum Marks; subject
    rows carry theory/practical/total.
    """
    got = _field_map(REAL_TN_HSC_OCR)
    assert "Theory" not in got, got
    assert "Practical" not in got, got
    assert got.get("Total Marks") == "0824", got
    assert got.get("Maximum Marks") == "200", got

    rows = parse_stacked_marksheet_subjects(REAL_TN_HSC_OCR)
    physics = next(s for s in rows if (s.get("name") or "").lower() == "physics")
    assert physics.get("theory") == "097", physics
    assert physics.get("practical") == "050", physics
    assert physics.get("total") == "147" or physics.get("marks") == "147", physics
    # No duplicate Total Marks subject rows.
    names = [(s.get("name") or "").lower() for s in rows]
    assert names.count("total marks") == 0
    assert "theory" not in names and "practical" not in names


def test_table_matrix_reconstructs_theory_practical() -> None:
    from app.services.scanx_ocr_blocks import (
        OcrTextBlock,
        reconstruct_table_matrices,
        format_table_matrices_as_text,
    )
    from app.services.scanx_academic_parse import parse_subjects_from_table_rows

    # Synthetic grid: header + Physics row (boxes on a 3-col layout).
    def _blk(i: int, text: str, x0: float, y0: float, w: float = 40.0, h: float = 14.0):
        return OcrTextBlock(
            block_id=f"b{i}",
            bounding_box=[[x0, y0], [x0 + w, y0], [x0 + w, y0 + h], [x0, y0 + h]],
            raw_ocr_text=text,
            cleaned_text=text,
            confidence=0.95,
            is_low_confidence=False,
            reading_order_index=i,
            page_index=0,
        )

    blocks = [
        _blk(0, "SUBJECT", 10, 10),
        _blk(1, "Theory", 80, 10),
        _blk(2, "Practical", 150, 10),
        _blk(3, "Total", 220, 10),
        _blk(4, "PHYSICS", 10, 40),
        _blk(5, "097", 80, 40),
        _blk(6, "050", 150, 40),
        _blk(7, "147", 220, 40),
        _blk(8, "CHEMISTRY", 10, 70),
        _blk(9, "060", 80, 70),
        _blk(10, "050", 150, 70),
        _blk(11, "110", 220, 70),
        _blk(12, "MATHEMATICS", 10, 100),
        _blk(13, "155", 80, 100),
        _blk(14, "000", 150, 100),
        _blk(15, "155", 220, 100),
    ]
    matrices = reconstruct_table_matrices(blocks, min_rows=3, min_cols=3)
    assert matrices, "expected reconstructed matrix"
    text = format_table_matrices_as_text(matrices)
    assert "Theory" in text and "Practical" in text
    rows = parse_subjects_from_table_rows(matrices[0])
    names = {(s.get("name") or "").lower() for s in rows}
    assert "physics" in names
    assert "mathematics" in names
    physics = next(s for s in rows if (s.get("name") or "").lower() == "physics")
    assert physics.get("marks") == "147"
    assert physics.get("theory") == "097", physics
    assert physics.get("practical") == "050", physics
    assert physics.get("total") == "147", physics


def test_docx_born_digital_theory_practical() -> None:
    """Native DOCX tables must extract text + theory/practical/total columns."""
    from io import BytesIO

    from docx import Document

    from app.services.scanx_embeddings import normalize_native_document_text
    from app.services.scanx_validation import extract_docx_text

    doc = Document()
    doc.add_paragraph("Name: Sample Student")
    doc.add_paragraph("Roll No: 99999")
    table = doc.add_table(rows=3, cols=4)
    for i, h in enumerate(["Subject", "Theory", "Practical", "Total Marks"]):
        table.rows[0].cells[i].text = h
    table.rows[1].cells[0].text = "Physics"
    table.rows[1].cells[1].text = "70"
    table.rows[1].cells[2].text = "50"
    table.rows[1].cells[3].text = "120"
    table.rows[2].cells[0].text = "Chemistry"
    table.rows[2].cells[1].text = "60"
    table.rows[2].cells[2].text = "40"
    table.rows[2].cells[3].text = "100"
    buf = BytesIO()
    doc.save(buf)
    joined, subjects = extract_docx_text(buf.getvalue())
    assert joined and "Physics" in joined and "Roll No" in joined
    assert normalize_native_document_text(joined)
    assert len(subjects) >= 2
    phys = next(s for s in subjects if s["name"].lower() == "physics")
    assert phys.get("theory") == "70"
    assert phys.get("practical") == "50"
    assert phys.get("total") == "120" or phys.get("marks") == "120"


def test_unicode_tamil_subject_name() -> None:
    """Local-script subject names must survive name_ok / subjects table."""
    text = "SUBJECT\nதமிழ்\n169\nONE SIX NINE\n(P)\nENGLISH\n137\n"
    rows = parse_subjects_from_text(text)
    assert any(
        "தமிழ்" in (s.get("name") or "") and str(s.get("marks") or "") == "169"
        for s in rows
    ), rows


if __name__ == "__main__":
    test_tamil_169_all_layouts()
    test_stacked_parser_real_ocr()
    test_resolve_merges_without_llm_when_stacked_ok()
    test_transcript_fields_dob_roll_group()
    test_transcript_fields_after_normalize()
    test_tmr_label_not_truncated_to_mr()
    test_mathematics_survives_misplaced_total()
    test_categorizer_includes_tamil_subjects()
    test_categorizer_real_ocr()
    test_mr_code_split_fragments_merge()
    test_unicode_tamil_subject_name()
    test_light_regex_clean_preserves_tmr()
    test_gibberish_ocr_artifacts_dropped()
    test_theory_practical_total_fields()
    test_table_matrix_reconstructs_theory_practical()
    test_docx_born_digital_theory_practical()
    print("ALL TAMIL / FIELD ASSERTIONS PASSED")
    # Optional live LLM smoke (skipped if Ollama down).
    try:
        from app.services.scanx_llm_subjects import extract_subjects_via_llm

        llm_rows = extract_subjects_via_llm(REAL_TN_HSC_OCR, timeout_seconds=60)
        print("LLM subjects:", llm_rows)
        if llm_rows:
            assert _has_tamil_169(llm_rows), llm_rows
            print("LLM TAMIL 169 OK")
    except Exception as exc:
        print("LLM smoke skipped/failed:", exc)
